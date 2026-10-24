import os
from datetime import datetime, date, timedelta
from functools import wraps
from io import BytesIO

from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import func

app = Flask(__name__)
# Always set SECRET_KEY in production (Cloud Run environment variable or Secret Manager).
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'local-development-only-change-before-deploy')
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_SECURE'] = os.environ.get('SESSION_COOKIE_SECURE', '').lower() in {'1', 'true', 'yes'}

# Cloud Run should use managed PostgreSQL because its container filesystem is ephemeral.
# DATABASE_URL accepts a PostgreSQL URL. SQLite remains available for local development.
database_url = os.environ.get('DATABASE_URL', '').strip()
if database_url:
    if database_url.startswith('postgres://'):
        database_url = 'postgresql://' + database_url[len('postgres://'):]
    if database_url.startswith('postgresql://'):
        database_url = database_url.replace('postgresql://', 'postgresql+psycopg://', 1)
    app.config['SQLALCHEMY_DATABASE_URI'] = database_url
    app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'pool_pre_ping': True, 'pool_recycle': 300}
else:
    instance_dir = os.path.join(app.root_path, 'instance')
    os.makedirs(instance_dir, exist_ok=True)
    database_path = os.path.abspath(os.environ.get('DB_PATH', os.path.join(instance_dir, 'supermart.db')))
    os.makedirs(os.path.dirname(database_path), exist_ok=True)
    app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///' + database_path
    app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'connect_args': {'check_same_thread': False}}
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

class Supermarket(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(160), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    is_superadmin = db.Column(db.Boolean, default=False, nullable=False)
    permissions = db.Column(db.String(500), default='sales,purchases,products,expenses,reports,accounting,receivables_payables', nullable=False)

class Product(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(150), nullable=False)
    sku = db.Column(db.String(60), nullable=True)
    unit = db.Column(db.String(20), default='pcs')
    cost_price = db.Column(db.Float, default=0)
    selling_price = db.Column(db.Float, default=0)
    stock_qty = db.Column(db.Float, default=0)
    low_stock_level = db.Column(db.Float, default=5)
    expiry_date = db.Column(db.Date, nullable=True, index=True)
    gst_rate = db.Column(db.Float, default=0, nullable=False)
    cst_rate = db.Column(db.Float, default=0, nullable=False)
    active = db.Column(db.Boolean, default=True)
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'name', name='uq_product_name_store'),)

class Purchase(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    supplier = db.Column(db.String(150), default='')
    quantity = db.Column(db.Float, nullable=False)
    unit_cost = db.Column(db.Float, nullable=False)
    total = db.Column(db.Float, nullable=False)
    purchase_date = db.Column(db.Date, default=date.today, index=True)
    expiry_date = db.Column(db.Date, nullable=True, index=True)
    note = db.Column(db.String(250), default='')
    paid = db.Column(db.Float, default=0, nullable=False)
    payment_method = db.Column(db.String(30), default='Credit', nullable=False)
    base_total = db.Column(db.Float, default=0, nullable=False)
    tax_type = db.Column(db.String(10), default='GST', nullable=False)
    supply_type = db.Column(db.String(20), default='Local', nullable=False)
    cgst = db.Column(db.Float, default=0, nullable=False)
    sgst = db.Column(db.Float, default=0, nullable=False)
    igst = db.Column(db.Float, default=0, nullable=False)
    cst = db.Column(db.Float, default=0, nullable=False)
    tax_amount = db.Column(db.Float, default=0, nullable=False)
    product = db.relationship('Product')

class Sale(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    invoice_no = db.Column(db.String(40), nullable=False)
    sale_date = db.Column(db.DateTime, default=datetime.now, index=True)
    customer_name = db.Column(db.String(150), default='Walk-in customer')
    payment_method = db.Column(db.String(30), default='Cash')
    subtotal = db.Column(db.Float, nullable=False)
    discount = db.Column(db.Float, default=0)
    total = db.Column(db.Float, nullable=False)
    paid = db.Column(db.Float, default=0)
    tax_type = db.Column(db.String(10), default='GST', nullable=False)
    supply_type = db.Column(db.String(20), default='Local', nullable=False)
    cgst = db.Column(db.Float, default=0, nullable=False)
    sgst = db.Column(db.Float, default=0, nullable=False)
    igst = db.Column(db.Float, default=0, nullable=False)
    cst = db.Column(db.Float, default=0, nullable=False)
    tax_amount = db.Column(db.Float, default=0, nullable=False)
    items = db.relationship('SaleItem', backref='sale', cascade='all, delete-orphan')
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'invoice_no', name='uq_invoice_store'),)

class SaleItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    sale_id = db.Column(db.Integer, db.ForeignKey('sale.id'), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    product_name = db.Column(db.String(150), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    unit_price = db.Column(db.Float, nullable=False)
    cost_price = db.Column(db.Float, nullable=False)
    line_total = db.Column(db.Float, nullable=False)
    product = db.relationship('Product')

class PaymentTransaction(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    direction = db.Column(db.String(10), nullable=False)  # receipt or payment
    sale_id = db.Column(db.Integer, db.ForeignKey('sale.id'), nullable=True)
    purchase_id = db.Column(db.Integer, db.ForeignKey('purchase.id'), nullable=True)
    party_name = db.Column(db.String(150), default='')
    amount = db.Column(db.Float, nullable=False)
    payment_method = db.Column(db.String(30), nullable=False, default='Cash')
    payment_date = db.Column(db.Date, default=date.today, nullable=False, index=True)
    reference = db.Column(db.String(100), default='')
    note = db.Column(db.String(250), default='')
    sale = db.relationship('Sale')
    purchase = db.relationship('Purchase')

class SalesReturn(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    sale_id = db.Column(db.Integer, db.ForeignKey('sale.id'), nullable=False, index=True)
    sale_item_id = db.Column(db.Integer, db.ForeignKey('sale_item.id'), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    unit_price = db.Column(db.Float, nullable=False)
    amount = db.Column(db.Float, nullable=False)
    return_date = db.Column(db.Date, default=date.today, nullable=False, index=True)
    reason = db.Column(db.String(250), default='')
    sale = db.relationship('Sale')
    sale_item = db.relationship('SaleItem')
    product = db.relationship('Product')

class PurchaseReturn(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    purchase_id = db.Column(db.Integer, db.ForeignKey('purchase.id'), nullable=False, index=True)
    product_id = db.Column(db.Integer, db.ForeignKey('product.id'), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    unit_cost = db.Column(db.Float, nullable=False)
    amount = db.Column(db.Float, nullable=False)
    return_date = db.Column(db.Date, default=date.today, nullable=False, index=True)
    reason = db.Column(db.String(250), default='')
    purchase = db.relationship('Purchase')
    product = db.relationship('Product')

class ExpenseCategory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(100), nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'name', name='uq_expense_category_store'),)

class PaymentMethod(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    name = db.Column(db.String(60), nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    __table_args__ = (db.UniqueConstraint('supermarket_id', 'name', name='uq_payment_method_store'),)

class Expense(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    supermarket_id = db.Column(db.Integer, db.ForeignKey('supermarket.id'), nullable=False, index=True)
    expense_date = db.Column(db.Date, default=date.today, index=True)
    category = db.Column(db.String(100), nullable=False)
    description = db.Column(db.String(250), default='')
    amount = db.Column(db.Float, nullable=False)
    payment_method = db.Column(db.String(30), default='Cash')

def seed_store_options(store_id_value):
    for category_name in ['Rent','Electricity','Water','Staff wages','Transport / delivery','Repairs & maintenance','Internet / phone','Packaging','Other']:
        if not ExpenseCategory.query.filter_by(supermarket_id=store_id_value, name=category_name).first():
            db.session.add(ExpenseCategory(supermarket_id=store_id_value, name=category_name))
    for method_name in ['Cash','UPI','Card','Bank transfer','Credit','Cheque']:
        if not PaymentMethod.query.filter_by(supermarket_id=store_id_value, name=method_name).first():
            db.session.add(PaymentMethod(supermarket_id=store_id_value, name=method_name))

with app.app_context():
    db.create_all()
    # Lightweight SQLite schema migration for existing installations.
    from sqlalchemy import inspect
    inspector = inspect(db.engine)
    purchase_columns = {column['name'] for column in inspector.get_columns('purchase')}
    product_columns = {column['name'] for column in inspector.get_columns('product')}
    purchase_columns = {column['name'] for column in inspector.get_columns('purchase')}
    sale_columns = {column['name'] for column in inspector.get_columns('sale')}
    with db.engine.begin() as connection:
        for col, ddl in {
            'gst_rate':'FLOAT NOT NULL DEFAULT 0', 'cst_rate':'FLOAT NOT NULL DEFAULT 0'
        }.items():
            if col not in product_columns: connection.exec_driver_sql(f'ALTER TABLE product ADD COLUMN {col} {ddl}')
        for col, ddl in {
            'base_total':'FLOAT NOT NULL DEFAULT 0','tax_type':"VARCHAR(10) NOT NULL DEFAULT 'GST'",'supply_type':"VARCHAR(20) NOT NULL DEFAULT 'Local'",'cgst':'FLOAT NOT NULL DEFAULT 0','sgst':'FLOAT NOT NULL DEFAULT 0','igst':'FLOAT NOT NULL DEFAULT 0','cst':'FLOAT NOT NULL DEFAULT 0','tax_amount':'FLOAT NOT NULL DEFAULT 0'
        }.items():
            if col not in purchase_columns: connection.exec_driver_sql(f'ALTER TABLE purchase ADD COLUMN {col} {ddl}')
            if col not in sale_columns: connection.exec_driver_sql(f'ALTER TABLE sale ADD COLUMN {col} {ddl}')
        connection.exec_driver_sql('UPDATE purchase SET base_total = total WHERE base_total = 0')
        if 'expiry_date' not in purchase_columns:
            connection.exec_driver_sql('ALTER TABLE purchase ADD COLUMN expiry_date DATE')
        if 'expiry_date' not in product_columns:
            connection.exec_driver_sql('ALTER TABLE product ADD COLUMN expiry_date DATE')
        if 'paid' not in purchase_columns:
            connection.exec_driver_sql('ALTER TABLE purchase ADD COLUMN paid FLOAT NOT NULL DEFAULT 0')
        if 'payment_method' not in purchase_columns:
            connection.exec_driver_sql("ALTER TABLE purchase ADD COLUMN payment_method VARCHAR(30) NOT NULL DEFAULT 'Credit'")
        user_columns = {column['name'] for column in inspector.get_columns('user')}
        if 'is_superadmin' not in user_columns:
            connection.exec_driver_sql('ALTER TABLE \"user\" ADD COLUMN is_superadmin BOOLEAN NOT NULL DEFAULT FALSE' if db.engine.dialect.name == 'postgresql' else 'ALTER TABLE \"user\" ADD COLUMN is_superadmin BOOLEAN NOT NULL DEFAULT 0')
        if 'permissions' not in user_columns:
            connection.exec_driver_sql("ALTER TABLE \"user\" ADD COLUMN permissions VARCHAR(500) NOT NULL DEFAULT 'sales,purchases,products,expenses,reports,accounting,receivables_payables'")
    # Provision a global administrator on first run. Override credentials using environment variables.
    admin_email = os.environ.get('SUPERADMIN_EMAIL', 'admin@supermart.local').strip().lower()
    admin_password = os.environ.get('SUPERADMIN_PASSWORD', 'SuperAdmin@2026')
    system_store = Supermarket.query.filter_by(name='SuperMart Administration').first()
    if not system_store:
        system_store = Supermarket(name='SuperMart Administration')
        db.session.add(system_store); db.session.flush()
    global_admin = User.query.filter_by(email=admin_email).first()
    if not global_admin:
        global_admin = User(supermarket_id=system_store.id, name='Global Administrator', email=admin_email,
                            password_hash=generate_password_hash(admin_password), is_superadmin=True, permissions='all')
        db.session.add(global_admin); db.session.commit()
    elif not global_admin.is_superadmin:
        global_admin.is_superadmin = True; global_admin.permissions = 'all'; db.session.commit()
    # Seed configurable expense categories and payment methods for each store.
    for store in Supermarket.query.all():
        seed_store_options(store.id)
    db.session.commit()

def store_id():
    return session.get('active_store_id') or session.get('supermarket_id')

PERMISSIONS = {'sales':'sales', 'sale_receipt':'sales', 'sales_history':'sales', 'purchases':'purchases', 'products':'products', 'expenses':'expenses', 'reports':'reports', 'export_reports':'reports', 'accounting':'accounting', 'receivables_payables':'receivables_payables', 'export_receivables_payables':'receivables_payables', 'receive_customer_payment':'receivables_payables', 'pay_supplier':'receivables_payables', 'payment_receipt':'receivables_payables', 'returns':'purchases', 'sales_return':'sales', 'purchase_return':'purchases', 'aging_report':'reports', 'export_aging_report':'reports', 'supplier_report':'reports', 'supplier_report_export':'reports', 'customer_report':'reports', 'customer_report_export':'reports', 'tax_report':'reports', 'tax_report_export':'reports', 'expiry_report':'products', 'expiry_report_export':'products', 'settings_lists':'expenses', 'toggle_setting_list':'expenses'}

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get('user_id'):
            return redirect(url_for('login'))
        current = db.session.get(User, session['user_id'])
        if not current:
            session.clear(); return redirect(url_for('login'))
        required = PERMISSIONS.get(fn.__name__)
        if required and not current.is_superadmin and required not in (current.permissions or '').split(','):
            flash('Your account does not have permission to access that section.', 'error')
            return redirect(url_for('dashboard'))
        return fn(*args, **kwargs)
    return wrapped

def money(value):
    return f"₹{(value or 0):,.2f}"
app.jinja_env.filters['money'] = money

def date_range():
    today = date.today()
    start_raw = request.args.get('start', '')
    end_raw = request.args.get('end', '')
    try: start = datetime.strptime(start_raw, '%Y-%m-%d').date() if start_raw else today.replace(day=1)
    except ValueError: start = today.replace(day=1)
    try: end = datetime.strptime(end_raw, '%Y-%m-%d').date() if end_raw else today
    except ValueError: end = today
    if end < start: start, end = end, start
    return start, end

@app.context_processor
def common_context():
    return {'store_name': session.get('store_name', 'SuperMart'), 'today': date.today().isoformat(), 'is_superadmin': bool(session.get('is_superadmin'))}

@app.route('/')
def index():
    if not session.get('user_id'): return redirect(url_for('login'))
    return redirect(url_for('dashboard'))

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        store = request.form.get('store_name','').strip()
        name = request.form.get('name','').strip()
        email = request.form.get('email','').strip().lower()
        password = request.form.get('password','')
        if not all([store, name, email, password]):
            flash('Please complete every field.', 'error')
        elif len(password) < 8:
            flash('Password must contain at least 8 characters.', 'error')
        elif User.query.filter_by(email=email).first():
            flash('This email is already registered. Please sign in.', 'error')
        else:
            supermarket = Supermarket(name=store)
            db.session.add(supermarket); db.session.flush(); seed_store_options(supermarket.id)
            user = User(supermarket_id=supermarket.id, name=name, email=email, password_hash=generate_password_hash(password))
            db.session.add(user); db.session.commit()
            session.update(user_id=user.id, supermarket_id=supermarket.id, active_store_id=supermarket.id, store_name=supermarket.name, user_name=user.name, is_superadmin=False)
            flash('Your supermarket workspace is ready.', 'success')
            return redirect(url_for('dashboard'))
    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email','').strip().lower()
        password = request.form.get('password','')
        user = User.query.filter_by(email=email).first()
        if user and check_password_hash(user.password_hash, password):
            supermarket = db.session.get(Supermarket, user.supermarket_id)
            session.clear(); session.update(user_id=user.id, supermarket_id=user.supermarket_id, active_store_id=(None if user.is_superadmin else user.supermarket_id), store_name=supermarket.name, user_name=user.name, is_superadmin=bool(user.is_superadmin))
            return redirect(url_for('admin_console') if user.is_superadmin else url_for('dashboard'))
        flash('Incorrect email or password.', 'error')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/admin')
@login_required
def admin_console():
    user = db.session.get(User, session['user_id'])
    if not user.is_superadmin:
        flash('Global administrator access required.', 'error'); return redirect(url_for('dashboard'))
    stores = Supermarket.query.order_by(Supermarket.name).all()
    users = User.query.filter_by(is_superadmin=False).order_by(User.email).all()
    summaries = []
    for store in stores:
        summaries.append({'store': store,
            'sales': db.session.query(func.coalesce(func.sum(Sale.total), 0)).filter_by(supermarket_id=store.id).scalar() or 0,
            'receivable': db.session.query(func.coalesce(func.sum(Sale.total - Sale.paid), 0)).filter(Sale.supermarket_id == store.id, Sale.total > Sale.paid).scalar() or 0,
            'payable': db.session.query(func.coalesce(func.sum(Purchase.total - Purchase.paid), 0)).filter(Purchase.supermarket_id == store.id, Purchase.total > Purchase.paid).scalar() or 0,
            'users': User.query.filter_by(supermarket_id=store.id, is_superadmin=False).count(),
            'products': Product.query.filter_by(supermarket_id=store.id).count()})
    return render_template('admin.html', stores=stores, users=users, summaries=summaries, store_names={store.id: store.name for store in stores})

@app.route('/admin/store/create', methods=['POST'])
@login_required
def admin_create_store():
    admin = db.session.get(User, session['user_id'])
    if not admin or not admin.is_superadmin: return ('Forbidden', 403)
    store_name = request.form.get('store_name', '').strip()
    owner_name = request.form.get('owner_name', '').strip()
    owner_email = request.form.get('owner_email', '').strip().lower()
    owner_password = request.form.get('owner_password', '')
    if not all([store_name, owner_name, owner_email, owner_password]) or len(owner_password) < 8:
        flash('Complete all fields. The owner password must be at least 8 characters.', 'error')
        return redirect(url_for('admin_console'))
    if User.query.filter_by(email=owner_email).first():
        flash('That email is already registered. Use a unique email.', 'error')
        return redirect(url_for('admin_console'))
    store = Supermarket(name=store_name)
    db.session.add(store); db.session.flush(); seed_store_options(store.id)
    owner = User(supermarket_id=store.id, name=owner_name, email=owner_email,
                 password_hash=generate_password_hash(owner_password))
    db.session.add(owner); db.session.commit()
    flash(f'Store {store_name} created with its own owner login.', 'success')
    return redirect(url_for('admin_console'))

@app.route('/admin/store/<int:sid>/rename', methods=['POST'])
@login_required
def admin_rename_store(sid):
    admin = db.session.get(User, session['user_id'])
    if not admin or not admin.is_superadmin: return ('Forbidden', 403)
    store = db.session.get(Supermarket, sid)
    if not store: return ('Store not found', 404)
    name = request.form.get('store_name', '').strip()
    if not name:
        flash('Store name cannot be blank.', 'error')
    else:
        store.name = name; db.session.commit()
        if session.get('active_store_id') == store.id: session['store_name'] = store.name
        flash('Store name updated.', 'success')
    return redirect(url_for('admin_console'))

@app.route('/admin/store/<int:sid>/open', methods=['POST'])
@login_required
def admin_open_store(sid):
    user = db.session.get(User, session['user_id'])
    if not user.is_superadmin: return ('Forbidden', 403)
    store = db.session.get(Supermarket, sid)
    if not store: return ('Store not found', 404)
    session['active_store_id'] = store.id; session['store_name'] = store.name
    flash(f'Now managing {store.name}. Use Global Admin to return to all stores.', 'success')
    return redirect(url_for('dashboard'))

@app.route('/admin/users/<int:user_id>/permissions', methods=['POST'])
@login_required
def admin_permissions(user_id):
    admin = db.session.get(User, session['user_id'])
    if not admin.is_superadmin: return ('Forbidden', 403)
    target = db.session.get(User, user_id)
    if not target or target.is_superadmin: return ('User not found', 404)
    allowed = {'sales','purchases','products','expenses','reports','accounting','receivables_payables'}
    selected = set(request.form.getlist('permissions')) & allowed
    target.permissions = ','.join(sorted(selected))
    db.session.commit(); flash(f'Permissions updated for {target.email}.', 'success')
    return redirect(url_for('admin_console'))

@app.route('/admin/all-stores')
@login_required
def admin_all_stores():
    user = db.session.get(User, session['user_id'])
    if not user.is_superadmin: return ('Forbidden', 403)
    session.pop('active_store_id', None); session['store_name'] = 'All Stores'
    return redirect(url_for('admin_console'))

@app.route('/dashboard')
@login_required
def dashboard():
    if session.get('is_superadmin') and not session.get('active_store_id'):
        return redirect(url_for('admin_console'))
    sid = store_id(); today = date.today()
    sales_today = db.session.query(func.coalesce(func.sum(Sale.total),0)).filter(Sale.supermarket_id==sid, func.date(Sale.sale_date)==today.isoformat()).scalar() or 0
    purchases_today = db.session.query(func.coalesce(func.sum(Purchase.total),0)).filter_by(supermarket_id=sid, purchase_date=today).scalar() or 0
    expenses_today = db.session.query(func.coalesce(func.sum(Expense.amount),0)).filter_by(supermarket_id=sid, expense_date=today).scalar() or 0
    product_count = Product.query.filter_by(supermarket_id=sid, active=True).count()
    low_stock = Product.query.filter(Product.supermarket_id==sid, Product.active.is_(True), Product.stock_qty <= Product.low_stock_level).order_by(Product.stock_qty.asc()).limit(8).all()
    recent_sales = Sale.query.filter_by(supermarket_id=sid).order_by(Sale.sale_date.desc()).limit(6).all()
    return render_template('dashboard.html', sales_today=sales_today, purchases_today=purchases_today, expenses_today=expenses_today, product_count=product_count, low_stock=low_stock, recent_sales=recent_sales)

@app.route('/products', methods=['GET','POST'])
@login_required
def products():
    sid=store_id()
    if request.method=='POST':
        name=request.form.get('name','').strip(); sku=request.form.get('sku','').strip()
        try:
            cost=float(request.form.get('cost_price',0)); price=float(request.form.get('selling_price',0)); qty=float(request.form.get('stock_qty',0)); low=float(request.form.get('low_stock_level',5)); gst_rate=float(request.form.get('gst_rate',0)); cst_rate=float(request.form.get('cst_rate',0))
            if not name or min(cost,price,qty,low,gst_rate,cst_rate)<0 or gst_rate>100 or cst_rate>100: raise ValueError()
        except ValueError:
            flash('Enter a product name and valid non-negative numbers.', 'error'); return redirect(url_for('products'))
        existing=Product.query.filter_by(supermarket_id=sid,name=name).first()
        if existing:
            flash('A product with this name already exists.', 'error')
        else:
            expiry_raw=request.form.get('expiry_date','').strip()
            expiry=datetime.strptime(expiry_raw,'%Y-%m-%d').date() if expiry_raw else None
            db.session.add(Product(supermarket_id=sid,name=name,sku=sku or None,unit=request.form.get('unit','pcs'),cost_price=cost,selling_price=price,stock_qty=qty,low_stock_level=low,expiry_date=expiry,gst_rate=gst_rate,cst_rate=cst_rate)); db.session.commit(); flash('Product added.', 'success')
        return redirect(url_for('products'))
    q=request.args.get('q','').strip()
    query=Product.query.filter_by(supermarket_id=sid,active=True)
    if q: query=query.filter(Product.name.ilike(f'%{q}%'))
    return render_template('products.html', products=query.order_by(Product.name).all(), q=q)

@app.route('/purchases', methods=['GET','POST'])
@login_required
def purchases():
    sid=store_id()
    if request.method=='POST':
        try:
            product=Product.query.filter_by(id=int(request.form['product_id']),supermarket_id=sid,active=True).first_or_404()
            qty=float(request.form['quantity']); cost=float(request.form['unit_cost'])
            if qty<=0 or cost<0: raise ValueError()
            pdate=datetime.strptime(request.form.get('purchase_date') or date.today().isoformat(),'%Y-%m-%d').date()
            base_total=round(qty*cost,2)
            tax_type=request.form.get('tax_type','GST').upper(); supply_type=request.form.get('supply_type','Local')
            if tax_type not in ('GST','CST') or supply_type not in ('Local','Inter-state'): raise ValueError()
            rate=float(product.gst_rate or 0) if tax_type=='GST' else float(product.cst_rate or 0)
            tax=round(base_total*rate/100,2); cgst=sgst=igst=cst=0.0
            if tax_type=='GST':
                if supply_type=='Local': cgst=round(tax/2,2); sgst=round(tax-cgst,2)
                else: igst=tax
            else: cst=tax
            total=round(base_total+tax,2)
            expiry_raw=request.form.get('expiry_date','').strip()
            expiry=datetime.strptime(expiry_raw,'%Y-%m-%d').date() if expiry_raw else None
            db.session.add(Purchase(supermarket_id=sid,product_id=product.id,supplier=request.form.get('supplier','').strip(),quantity=qty,unit_cost=cost,total=total,base_total=base_total,tax_type=tax_type,supply_type=supply_type,cgst=cgst,sgst=sgst,igst=igst,cst=cst,tax_amount=tax,purchase_date=pdate,expiry_date=expiry,note=request.form.get('note','').strip(),paid=min(max(float(request.form.get('paid',0)),0),total),payment_method=request.form.get('payment_method','Credit')))
            product.stock_qty += qty; product.cost_price=cost
            if expiry: product.expiry_date=expiry
            db.session.commit(); flash('Purchase saved and stock updated.', 'success')
        except (ValueError, KeyError):
            db.session.rollback(); flash('Please enter valid purchase details.', 'error')
        return redirect(url_for('purchases'))
    items=Purchase.query.filter_by(supermarket_id=sid).order_by(Purchase.purchase_date.desc(),Purchase.id.desc()).limit(100).all()
    return render_template('purchases.html', products=Product.query.filter_by(supermarket_id=sid,active=True).order_by(Product.name).all(), purchases=items, payment_methods=PaymentMethod.query.filter_by(supermarket_id=sid,active=True).order_by(PaymentMethod.name).all())

@app.route('/sales', methods=['GET','POST'])
@login_required
def sales():
    sid=store_id()
    products_list=Product.query.filter_by(supermarket_id=sid,active=True).order_by(Product.name).all()
    if request.method=='POST':
        try:
            ids=request.form.getlist('product_id[]'); qtys=request.form.getlist('quantity[]')
            discount=float(request.form.get('discount',0)); paid=float(request.form.get('paid',0))
            if discount<0 or paid<0: raise ValueError()
            line_data=[]; subtotal=0; requested_by_product={}; gst_tax=0.0; cst_tax=0.0; cgst=sgst=igst=cst=0.0
            tax_type=request.form.get('tax_type','GST').upper(); supply_type=request.form.get('supply_type','Local')
            if tax_type not in ('GST','CST') or supply_type not in ('Local','Inter-state'): raise ValueError('Choose a valid tax type and supply type.')
            for pid, qraw in zip(ids,qtys):
                if not pid or not qraw: continue
                product=Product.query.filter_by(id=int(pid),supermarket_id=sid,active=True).first()
                qty=float(qraw)
                if not product or qty<=0: raise ValueError('Invalid product or quantity.')
                requested_by_product[product.id] = requested_by_product.get(product.id, 0) + qty
                if product.stock_qty < requested_by_product[product.id]: raise ValueError(f'Not enough stock for {product.name}. Available: {product.stock_qty:g}')
                line_total=round(qty*product.selling_price,2); subtotal+=line_total
                rate=float(product.gst_rate or 0) if tax_type=='GST' else float(product.cst_rate or 0)
                line_tax=round(line_total*rate/100,2)
                if tax_type=='GST': gst_tax+=line_tax
                else: cst_tax+=line_tax
                line_data.append((product,qty,line_total,line_tax,rate))
            if not line_data: raise ValueError('Add at least one product.')
            if discount>subtotal: raise ValueError('Discount cannot exceed subtotal.')
            base_after_discount=subtotal-discount
            # Allocate discount proportionally before tax, by applying the transaction discount to the aggregate tax base.
            tax_total=round((gst_tax if tax_type=='GST' else cst_tax) * (base_after_discount/subtotal if subtotal else 0),2)
            if tax_type=='GST':
                if supply_type=='Local': cgst=round(tax_total/2,2); sgst=round(tax_total-cgst,2); igst=0.0
                else: igst=tax_total; cgst=sgst=0.0
                cst=0.0
            else: cst=tax_total; cgst=sgst=igst=0.0
            total=round(base_after_discount+tax_total,2)
            invoice='SM'+datetime.now().strftime('%y%m%d%H%M%S')+str(session['user_id'])
            sale=Sale(supermarket_id=sid,invoice_no=invoice,customer_name=request.form.get('customer_name','').strip() or 'Walk-in customer',payment_method=request.form.get('payment_method','Cash'),subtotal=subtotal,discount=discount,total=total,paid=min(paid,total),tax_type=tax_type,supply_type=supply_type,cgst=cgst,sgst=sgst,igst=igst,cst=cst,tax_amount=tax_total)
            db.session.add(sale); db.session.flush()
            for product,qty,line_total,line_tax,line_rate in line_data:
                db.session.add(SaleItem(sale_id=sale.id,product_id=product.id,product_name=product.name,quantity=qty,unit_price=product.selling_price,cost_price=product.cost_price,line_total=line_total))
                product.stock_qty-=qty
            db.session.commit(); return redirect(url_for('sale_receipt',sale_id=sale.id))
        except (ValueError, KeyError) as exc:
            db.session.rollback(); flash(str(exc) or 'Please enter valid sale details.', 'error')
            return redirect(url_for('sales'))
    return render_template('sales.html', products=products_list, payment_methods=PaymentMethod.query.filter_by(supermarket_id=sid,active=True).order_by(PaymentMethod.name).all())

@app.route('/sales/<int:sale_id>/receipt')
@login_required
def sale_receipt(sale_id):
    sale=Sale.query.filter_by(id=sale_id,supermarket_id=store_id()).first_or_404()
    return render_template('receipt.html', sale=sale)

@app.route('/sales/history')
@login_required
def sales_history():
    start,end=date_range(); sid=store_id()
    rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).order_by(Sale.sale_date.desc()).all()
    return render_template('sales_history.html', sales=rows,start=start.isoformat(),end=end.isoformat(),total=sum(x.total for x in rows))

@app.route('/expenses', methods=['GET','POST'])
@login_required
def expenses():
    sid=store_id()
    if request.method=='POST':
        try:
            amount=float(request.form.get('amount',0)); edate=datetime.strptime(request.form.get('expense_date') or date.today().isoformat(),'%Y-%m-%d').date()
            category=request.form.get('category','').strip(); method=request.form.get('payment_method','Cash').strip()
            valid_categories={x.name for x in ExpenseCategory.query.filter_by(supermarket_id=sid,active=True).all()}
            valid_methods={x.name for x in PaymentMethod.query.filter_by(supermarket_id=sid,active=True).all()}
            if amount<=0 or category not in valid_categories or method not in valid_methods: raise ValueError()
            db.session.add(Expense(supermarket_id=sid,expense_date=edate,category=category,description=request.form.get('description','').strip(),amount=amount,payment_method=method)); db.session.commit(); flash('Expense saved.', 'success')
        except ValueError:
            db.session.rollback(); flash('Enter a category and a valid amount.', 'error')
        return redirect(url_for('expenses'))
    rows=Expense.query.filter_by(supermarket_id=sid).order_by(Expense.expense_date.desc(),Expense.id.desc()).limit(100).all()
    return render_template('expenses.html',expenses=rows,categories=ExpenseCategory.query.filter_by(supermarket_id=sid,active=True).order_by(ExpenseCategory.name).all(),payment_methods=PaymentMethod.query.filter_by(supermarket_id=sid,active=True).order_by(PaymentMethod.name).all())

@app.route('/reports')
@login_required
def reports():
    start,end=date_range(); sid=store_id()
    sales_rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).order_by(Sale.sale_date.desc()).all()
    purchase_rows=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).order_by(Purchase.purchase_date.desc()).all()
    expense_rows=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).order_by(Expense.expense_date.desc()).all()
    cogs=sum(item.quantity*item.cost_price for sale in sales_rows for item in sale.items)
    sales_total=sum(s.total for s in sales_rows)-sum((db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(sale_id=s.id,supermarket_id=sid).scalar() or 0) for s in sales_rows); purchase_total=sum(p.total for p in purchase_rows)-sum((db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(purchase_id=p.id,supermarket_id=sid).scalar() or 0) for p in purchase_rows); expense_total=sum(e.amount for e in expense_rows)
    sales_returns=SalesReturn.query.filter(SalesReturn.supermarket_id==sid,SalesReturn.return_date>=start,SalesReturn.return_date<=end).order_by(SalesReturn.return_date.desc()).all()
    purchase_returns=PurchaseReturn.query.filter(PurchaseReturn.supermarket_id==sid,PurchaseReturn.return_date>=start,PurchaseReturn.return_date<=end).order_by(PurchaseReturn.return_date.desc()).all()
    return render_template('reports.html',start=start.isoformat(),end=end.isoformat(),sales=sales_rows,purchases=purchase_rows,expenses=expense_rows,sales_returns=sales_returns,purchase_returns=purchase_returns,sales_total=sales_total,purchase_total=purchase_total,expense_total=expense_total,cogs=cogs,gross_profit=sales_total-cogs,net_estimate=sales_total-cogs-expense_total)

@app.route('/reports/export.xlsx')
@login_required
def export_reports():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    start,end=date_range(); sid=store_id(); wb=Workbook(); ws=wb.active; ws.title='Summary'
    ws.append([session.get('store_name','SuperMart'),'Report summary']); ws.append(['Period',f'{start} to {end}'])
    sales_rows=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all()
    purchase_rows=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    expense_rows=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).all()
    sales_total=sum(s.total for s in sales_rows)-sum((db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(sale_id=s.id,supermarket_id=sid).scalar() or 0) for s in sales_rows); purchase_total=sum(p.total for p in purchase_rows)-sum((db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(purchase_id=p.id,supermarket_id=sid).scalar() or 0) for p in purchase_rows); expense_total=sum(e.amount for e in expense_rows)
    cogs=sum(i.quantity*i.cost_price for s in sales_rows for i in s.items)
    for row in [['Sales',sales_total],['Purchases',purchase_total],['Other expenses',expense_total],['Cost of goods sold (estimated)',cogs],['Gross profit (estimated)',sales_total-cogs],['Net estimate after expenses',sales_total-cogs-expense_total]]: ws.append(row)
    ws2=wb.create_sheet('Sales'); ws2.append(['Invoice','Date','Customer','Payment','Subtotal','Discount','Total','Paid'])
    for s in sales_rows: ws2.append([s.invoice_no,s.sale_date.strftime('%Y-%m-%d %H:%M'),s.customer_name,s.payment_method,s.subtotal,s.discount,s.total,s.paid])
    ws3=wb.create_sheet('Purchases'); ws3.append(['Date','Product','Supplier','Quantity','Unit cost','Total','Paid','Outstanding','Payment method','Note'])
    for p in purchase_rows: ws3.append([p.purchase_date.isoformat(),p.product.name,p.supplier,p.quantity,p.unit_cost,p.total,p.paid or 0,max(0,p.total-(p.paid or 0)),p.payment_method,p.note])
    ws4=wb.create_sheet('Expenses'); ws4.append(['Date','Category','Description','Payment','Amount'])
    for e in expense_rows: ws4.append([e.expense_date.isoformat(),e.category,e.description,e.payment_method,e.amount])
    sr=wb.create_sheet('Sales Returns'); sr.append(['Return Date','Invoice','Customer','Product','Quantity','Amount','Reason'])
    for r in SalesReturn.query.filter(SalesReturn.supermarket_id==sid,SalesReturn.return_date>=start,SalesReturn.return_date<=end).all(): sr.append([r.return_date.isoformat(),r.sale.invoice_no,r.sale.customer_name,r.product.name,r.quantity,r.amount,r.reason])
    pr=wb.create_sheet('Purchase Returns'); pr.append(['Return Date','Supplier','Product','Quantity','Amount','Reason'])
    for r in PurchaseReturn.query.filter(PurchaseReturn.supermarket_id==sid,PurchaseReturn.return_date>=start,PurchaseReturn.return_date<=end).all(): pr.append([r.return_date.isoformat(),r.purchase.supplier,r.product.name,r.quantity,r.amount,r.reason])
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='203B53')
        for col in sheet.columns:
            width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
            sheet.column_dimensions[get_column_letter(col[0].column)].width=width
    output=BytesIO(); wb.save(output); output.seek(0)
    return send_file(output,as_attachment=True,download_name=f'supermart_report_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/tax-report')
@login_required
def tax_report():
    start,end=date_range(); sid=store_id()
    sales=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all()
    purchases=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    sales_returns=SalesReturn.query.filter(SalesReturn.supermarket_id==sid,SalesReturn.return_date>=start,SalesReturn.return_date<=end).all()
    purchase_returns=PurchaseReturn.query.filter(PurchaseReturn.supermarket_id==sid,PurchaseReturn.return_date>=start,PurchaseReturn.return_date<=end).all()
    # Return tax is estimated proportionally from the source transaction's effective tax rate.
    out={k:0.0 for k in ('sales_cgst','sales_sgst','sales_igst','sales_cst','purchase_cgst','purchase_sgst','purchase_igst','purchase_cst','sales_return_tax','purchase_return_tax')}
    for x in sales: out['sales_cgst']+=x.cgst or 0; out['sales_sgst']+=x.sgst or 0; out['sales_igst']+=x.igst or 0; out['sales_cst']+=x.cst or 0
    for x in purchases: out['purchase_cgst']+=x.cgst or 0; out['purchase_sgst']+=x.sgst or 0; out['purchase_igst']+=x.igst or 0; out['purchase_cst']+=x.cst or 0
    for r in sales_returns:
        src=r.sale; denom=max(0.01,src.subtotal-(src.discount or 0)); out['sales_return_tax'] += (src.tax_amount or 0)*(r.amount/denom)
    for r in purchase_returns:
        src=r.purchase; denom=max(0.01,src.base_total or src.total); out['purchase_return_tax'] += (src.tax_amount or 0)*(r.amount/denom)
    out={k:round(v,2) for k,v in out.items()}
    sales_gst=out['sales_cgst']+out['sales_sgst']+out['sales_igst']; purchase_gst=out['purchase_cgst']+out['purchase_sgst']+out['purchase_igst']
    net_gst=round(sales_gst-out['sales_return_tax']-purchase_gst+out['purchase_return_tax'],2)
    net_cst=round(out['sales_cst']-out['purchase_cst'],2)
    return render_template('tax_report.html',start=start.isoformat(),end=end.isoformat(),sales=sales,purchases=purchases,out=out,sales_gst=round(sales_gst,2),purchase_gst=round(purchase_gst,2),net_gst=net_gst,net_cst=net_cst,sales_returns=sales_returns,purchase_returns=purchase_returns)

@app.route('/tax-report/export.xlsx')
@login_required
def tax_report_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    start,end=date_range(); sid=store_id(); wb=Workbook(); ws=wb.active; ws.title='Tax Summary'
    ws.append(['Tax submission summary',session.get('store_name','SuperMart')]); ws.append(['Period',f'{start} to {end}'])
    sales=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all(); purchases=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    sales_returns=SalesReturn.query.filter(SalesReturn.supermarket_id==sid,SalesReturn.return_date>=start,SalesReturn.return_date<=end).all()
    purchase_returns=PurchaseReturn.query.filter(PurchaseReturn.supermarket_id==sid,PurchaseReturn.return_date>=start,PurchaseReturn.return_date<=end).all()
    sales_return_tax=sum((r.sale.tax_amount or 0)*(r.amount/max(0.01,r.sale.subtotal-(r.sale.discount or 0))) for r in sales_returns)
    purchase_return_tax=sum((r.purchase.tax_amount or 0)*(r.amount/max(0.01,r.purchase.base_total or r.purchase.total)) for r in purchase_returns)
    output_gst=sum((x.cgst or 0)+(x.sgst or 0)+(x.igst or 0) for x in sales)
    input_gst=sum((x.cgst or 0)+(x.sgst or 0)+(x.igst or 0) for x in purchases)
    for label,value in [('Sales taxable value (before tax)',sum(max(0,s.subtotal-(s.discount or 0)) for s in sales)),('Output CGST',sum(s.cgst or 0 for s in sales)),('Output SGST',sum(s.sgst or 0 for s in sales)),('Output IGST',sum(s.igst or 0 for s in sales)),('Output CST (legacy)',sum(s.cst or 0 for s in sales)),('Purchase taxable value (before tax)',sum(p.base_total or p.total for p in purchases)),('Input CGST',sum(p.cgst or 0 for p in purchases)),('Input SGST',sum(p.sgst or 0 for p in purchases)),('Input IGST',sum(p.igst or 0 for p in purchases)),('Input CST (legacy)',sum(p.cst or 0 for p in purchases)),('Estimated tax reduction from sales returns',sales_return_tax),('Estimated input tax reduction from purchase returns',purchase_return_tax),('Estimated net GST after returns',output_gst-sales_return_tax-input_gst+purchase_return_tax)]: ws.append([label,round(value,2)])
    detail=wb.create_sheet('Sales invoices'); detail.append(['Invoice','Date','Customer','Tax type','Supply type','Taxable value','CGST','SGST','IGST','CST','Tax','Invoice total'])
    for x in sales: detail.append([x.invoice_no,x.sale_date.isoformat(),x.customer_name,x.tax_type,x.supply_type,round(x.subtotal-(x.discount or 0),2),x.cgst,x.sgst,x.igst,x.cst,x.tax_amount,x.total])
    detail=wb.create_sheet('Purchase invoices'); detail.append(['Date','Supplier','Product','Tax type','Supply type','Taxable value','CGST','SGST','IGST','CST','Tax','Invoice total'])
    for x in purchases: detail.append([x.purchase_date.isoformat(),x.supplier,x.product.name,x.tax_type,x.supply_type,x.base_total,x.cgst,x.sgst,x.igst,x.cst,x.tax_amount,x.total])
    for sh in wb.worksheets:
        sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
        for c in sh[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='203B53')
        for col in sh.columns: sh.column_dimensions[col[0].column_letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
    output=BytesIO(); wb.save(output); output.seek(0); return send_file(output,as_attachment=True,download_name=f'tax_submission_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/accounting')
@login_required
def accounting():
    start,end=date_range(); sid=store_id()
    ss=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start.isoformat(),func.date(Sale.sale_date)<=end.isoformat()).all()
    ps=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    es=Expense.query.filter(Expense.supermarket_id==sid,Expense.expense_date>=start,Expense.expense_date<=end).all()
    sales_total=sum(s.total for s in ss)-sum((db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(sale_id=s.id,supermarket_id=sid).scalar() or 0) for s in ss); cogs=sum(i.quantity*i.cost_price for s in ss for i in s.items); purchase_total=sum(p.total for p in ps)-sum((db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(purchase_id=p.id,supermarket_id=sid).scalar() or 0) for p in ps); expense_total=sum(e.amount for e in es)
    return render_template('accounting.html',start=start.isoformat(),end=end.isoformat(),sales_total=sales_total,purchase_total=purchase_total,expense_total=expense_total,cogs=cogs,gross_profit=sales_total-cogs,net_estimate=sales_total-cogs-expense_total)


@app.route('/receivables-payables')
@login_required
def receivables_payables():
    sid = store_id()
    start, end = date_range()
    # Outstanding balances are shown across ALL dates so old unpaid invoices are never hidden.
    sales_rows = Sale.query.filter(Sale.supermarket_id == sid).order_by(Sale.sale_date.desc()).all()
    purchase_rows = Purchase.query.filter(Purchase.supermarket_id == sid).order_by(Purchase.purchase_date.desc()).all()
    for sale in sales_rows:
        sale.return_total = db.session.query(func.coalesce(func.sum(SalesReturn.amount), 0)).filter_by(sale_id=sale.id, supermarket_id=sid).scalar() or 0
    for purchase in purchase_rows:
        purchase.return_total = db.session.query(func.coalesce(func.sum(PurchaseReturn.amount), 0)).filter_by(purchase_id=purchase.id, supermarket_id=sid).scalar() or 0
    receivables = [sale for sale in sales_rows if sale.total - (sale.paid or 0) - sale.return_total > 0.009]
    payables = [purchase for purchase in purchase_rows if purchase.total - (purchase.paid or 0) - purchase.return_total > 0.009]
    transactions = PaymentTransaction.query.filter_by(supermarket_id=sid).filter(PaymentTransaction.payment_date >= start, PaymentTransaction.payment_date <= end).order_by(PaymentTransaction.payment_date.desc(), PaymentTransaction.id.desc()).all()
    return render_template('receivables_payables.html', start=start.isoformat(), end=end.isoformat(), receivables=receivables, payables=payables, transactions=transactions, payment_methods=PaymentMethod.query.filter_by(supermarket_id=sid,active=True).order_by(PaymentMethod.name).all(),
        receivable_total=sum(max(0, s.total - (s.paid or 0) - s.return_total) for s in receivables), payable_total=sum(max(0, p.total - (p.paid or 0) - p.return_total) for p in payables))

@app.route('/receivables-payables/export.xlsx')
@login_required
def export_receivables_payables():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    sid = store_id(); start, end = date_range(); wb = Workbook()
    ar = wb.active; ar.title = 'Accounts Receivable'
    ar.append(['Invoice', 'Date', 'Customer', 'Payment method', 'Invoice total', 'Amount received', 'Returns / credit', 'Outstanding'])
    sales_rows = Sale.query.filter(Sale.supermarket_id == sid).order_by(Sale.sale_date.desc()).all()
    for sale in sales_rows:
        returned = db.session.query(func.coalesce(func.sum(SalesReturn.amount), 0)).filter_by(sale_id=sale.id, supermarket_id=sid).scalar() or 0
        due = round(max(0, sale.total - (sale.paid or 0) - returned), 2)
        if due > 0.009: ar.append([sale.invoice_no, sale.sale_date.strftime('%Y-%m-%d %H:%M'), sale.customer_name, sale.payment_method, sale.total, sale.paid or 0, returned, due])
    ap = wb.create_sheet('Accounts Payable')
    ap.append(['Date', 'Supplier', 'Product', 'Payment method', 'Purchase total', 'Amount paid', 'Returns / credit', 'Outstanding'])
    purchase_rows = Purchase.query.filter(Purchase.supermarket_id == sid).order_by(Purchase.purchase_date.desc()).all()
    for purchase in purchase_rows:
        returned = db.session.query(func.coalesce(func.sum(PurchaseReturn.amount), 0)).filter_by(purchase_id=purchase.id, supermarket_id=sid).scalar() or 0
        due = round(max(0, purchase.total - (purchase.paid or 0) - returned), 2)
        if due > 0.009: ap.append([purchase.purchase_date.isoformat(), purchase.supplier, purchase.product.name, purchase.payment_method, purchase.total, purchase.paid or 0, returned, due])
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='203B53')
        for col in sheet.columns:
            width = min(max(max(len(str(c.value or '')) for c in col) + 2, 12), 32)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    output = BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=f'supermart_ar_ap_{start}_{end}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/payments/receive', methods=['POST'])
@login_required
def receive_customer_payment():
    sid = store_id()
    try:
        sale = Sale.query.filter_by(id=int(request.form['sale_id']), supermarket_id=sid).first_or_404()
        amount = round(float(request.form.get('amount', 0)), 2)
        if amount <= 0 or amount > round(sale.total - (sale.paid or 0) - (db.session.query(func.coalesce(func.sum(SalesReturn.amount), 0)).filter_by(sale_id=sale.id, supermarket_id=sid).scalar() or 0), 2) + 0.001:
            raise ValueError('Enter an amount greater than zero and no more than the outstanding balance.')
        method = request.form.get('payment_method', 'Cash')
        pdate = datetime.strptime(request.form.get('payment_date') or date.today().isoformat(), '%Y-%m-%d').date()
        sale.paid = round((sale.paid or 0) + amount, 2)
        tx = PaymentTransaction(supermarket_id=sid, direction='receipt', sale_id=sale.id, party_name=sale.customer_name, amount=amount, payment_method=method, payment_date=pdate, reference=request.form.get('reference','').strip(), note=request.form.get('note','').strip())
        db.session.add(tx); db.session.commit(); flash('Customer payment received and balance updated.', 'success')
    except (ValueError, KeyError) as exc:
        db.session.rollback(); flash(str(exc) or 'Invalid payment details.', 'error')
    return redirect(url_for('payment_receipt', payment_id=tx.id)) if 'tx' in locals() else redirect(url_for('receivables_payables'))

@app.route('/payments/pay-supplier', methods=['POST'])
@login_required
def pay_supplier():
    sid = store_id()
    try:
        purchase = Purchase.query.filter_by(id=int(request.form['purchase_id']), supermarket_id=sid).first_or_404()
        amount = round(float(request.form.get('amount', 0)), 2)
        if amount <= 0 or amount > round(purchase.total - (purchase.paid or 0) - (db.session.query(func.coalesce(func.sum(PurchaseReturn.amount), 0)).filter_by(purchase_id=purchase.id, supermarket_id=sid).scalar() or 0), 2) + 0.001:
            raise ValueError('Enter an amount greater than zero and no more than the outstanding balance.')
        method = request.form.get('payment_method', 'Cash')
        pdate = datetime.strptime(request.form.get('payment_date') or date.today().isoformat(), '%Y-%m-%d').date()
        purchase.paid = round((purchase.paid or 0) + amount, 2)
        tx = PaymentTransaction(supermarket_id=sid, direction='payment', purchase_id=purchase.id, party_name=purchase.supplier, amount=amount, payment_method=method, payment_date=pdate, reference=request.form.get('reference','').strip(), note=request.form.get('note','').strip())
        db.session.add(tx); db.session.commit(); flash('Supplier payment saved and payable balance updated.', 'success')
    except (ValueError, KeyError) as exc:
        db.session.rollback(); flash(str(exc) or 'Invalid payment details.', 'error')
    return redirect(url_for('payment_receipt', payment_id=tx.id)) if 'tx' in locals() else redirect(url_for('receivables_payables'))

@app.route('/payments/<int:payment_id>/receipt')
@login_required
def payment_receipt(payment_id):
    payment = PaymentTransaction.query.filter_by(id=payment_id, supermarket_id=store_id()).first_or_404()
    return render_template('payment_receipt.html', payment=payment)


@app.route('/returns', methods=['GET'])
@login_required
def returns():
    sid = store_id()
    sales_rows = Sale.query.filter_by(supermarket_id=sid).order_by(Sale.sale_date.desc()).limit(250).all()
    purchase_rows = Purchase.query.filter_by(supermarket_id=sid).order_by(Purchase.purchase_date.desc(), Purchase.id.desc()).limit(250).all()
    sale_items = SaleItem.query.join(Sale, SaleItem.sale_id == Sale.id).filter(Sale.supermarket_id == sid).order_by(Sale.sale_date.desc()).limit(500).all()
    sr = SalesReturn.query.filter_by(supermarket_id=sid).order_by(SalesReturn.return_date.desc(), SalesReturn.id.desc()).limit(100).all()
    pr = PurchaseReturn.query.filter_by(supermarket_id=sid).order_by(PurchaseReturn.return_date.desc(), PurchaseReturn.id.desc()).limit(100).all()
    sale_returned_qty = {i.id: (db.session.query(func.coalesce(func.sum(SalesReturn.quantity),0)).filter_by(sale_item_id=i.id, supermarket_id=sid).scalar() or 0) for i in sale_items}
    purchase_returned_qty = {i.id: (db.session.query(func.coalesce(func.sum(PurchaseReturn.quantity),0)).filter_by(purchase_id=i.id, supermarket_id=sid).scalar() or 0) for i in purchase_rows}
    return render_template('returns.html', sales=sales_rows, purchases=purchase_rows, sale_items=sale_items, sales_returns=sr, purchase_returns=pr, sale_returned_qty=sale_returned_qty, purchase_returned_qty=purchase_returned_qty)

@app.route('/returns/sales', methods=['POST'])
@login_required
def sales_return():
    sid = store_id()
    try:
        item = SaleItem.query.join(Sale, SaleItem.sale_id == Sale.id).filter(SaleItem.id == int(request.form['sale_item_id']), Sale.supermarket_id == sid).first_or_404()
        already = db.session.query(func.coalesce(func.sum(SalesReturn.quantity),0)).filter_by(sale_item_id=item.id, supermarket_id=sid).scalar() or 0
        qty = float(request.form.get('quantity', 0))
        if qty <= 0 or qty > item.quantity - already + 0.0001: raise ValueError(f'Return quantity must be above zero and no more than {item.quantity-already:g}.')
        product = db.session.get(Product, item.product_id)
        if not product or product.supermarket_id != sid: raise ValueError('Product not found in this store.')
        rdate = datetime.strptime(request.form.get('return_date') or date.today().isoformat(), '%Y-%m-%d').date()
        amount = round(qty * item.unit_price, 2)
        record = SalesReturn(supermarket_id=sid, sale_id=item.sale_id, sale_item_id=item.id, product_id=item.product_id, quantity=qty, unit_price=item.unit_price, amount=amount, return_date=rdate, reason=request.form.get('reason','').strip())
        product.stock_qty += qty
        db.session.add(record); db.session.commit()
        flash(f'Sales return saved ({money(amount)}). Stock increased. Any unpaid customer balance is reduced by the return value.', 'success')
    except (ValueError, KeyError) as exc:
        db.session.rollback(); flash(str(exc) or 'Invalid sales return.', 'error')
    return redirect(url_for('returns'))

@app.route('/returns/purchases', methods=['POST'])
@login_required
def purchase_return():
    sid = store_id()
    try:
        purchase = Purchase.query.filter_by(id=int(request.form['purchase_id']), supermarket_id=sid).first_or_404()
        already = db.session.query(func.coalesce(func.sum(PurchaseReturn.quantity),0)).filter_by(purchase_id=purchase.id, supermarket_id=sid).scalar() or 0
        qty = float(request.form.get('quantity', 0))
        if qty <= 0 or qty > purchase.quantity - already + 0.0001: raise ValueError(f'Return quantity must be above zero and no more than {purchase.quantity-already:g}.')
        product = db.session.get(Product, purchase.product_id)
        if not product or product.supermarket_id != sid: raise ValueError('Product not found in this store.')
        if qty > product.stock_qty + 0.0001: raise ValueError(f'Cannot return {qty:g}; current stock is only {product.stock_qty:g}.')
        rdate = datetime.strptime(request.form.get('return_date') or date.today().isoformat(), '%Y-%m-%d').date()
        amount = round(qty * purchase.unit_cost, 2)
        record = PurchaseReturn(supermarket_id=sid, purchase_id=purchase.id, product_id=purchase.product_id, quantity=qty, unit_cost=purchase.unit_cost, amount=amount, return_date=rdate, reason=request.form.get('reason','').strip())
        product.stock_qty -= qty
        db.session.add(record); db.session.commit()
        flash(f'Purchase return saved ({money(amount)}). Stock reduced. Any unpaid supplier balance is reduced by the return value.', 'success')
    except (ValueError, KeyError) as exc:
        db.session.rollback(); flash(str(exc) or 'Invalid purchase return.', 'error')
    return redirect(url_for('returns'))

@app.route('/reports/aging')
@login_required
def aging_report():
    sid = store_id(); today = date.today()
    sales = Sale.query.filter_by(supermarket_id=sid).all()
    purchases = Purchase.query.filter_by(supermarket_id=sid).all()
    ar_rows=[]; ap_rows=[]
    def buckets(age, due):
        return [due if age <= 0 else 0, due if 1 <= age <= 30 else 0, due if 31 <= age <= 60 else 0, due if 61 <= age <= 90 else 0, due if age > 90 else 0]
    for sale in sales:
        returned = db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(sale_id=sale.id, supermarket_id=sid).scalar() or 0
        due = round(max(0, sale.total - (sale.paid or 0) - returned),2)
        if due > .009:
            age=(today-sale.sale_date.date()).days; ar_rows.append({'ref':sale.invoice_no,'date':sale.sale_date.date(),'party':sale.customer_name,'total':sale.total,'paid':sale.paid or 0,'returns':returned,'due':due,'age':age,'buckets':buckets(age,due)})
    for purchase in purchases:
        returned = db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(purchase_id=purchase.id, supermarket_id=sid).scalar() or 0
        due=round(max(0,purchase.total-(purchase.paid or 0)-returned),2)
        if due > .009:
            age=(today-purchase.purchase_date).days; ap_rows.append({'ref':f'PUR-{purchase.id}','date':purchase.purchase_date,'party':purchase.supplier or 'Supplier not specified','total':purchase.total,'paid':purchase.paid or 0,'returns':returned,'due':due,'age':age,'buckets':buckets(age,due)})
    return render_template('aging_report.html', ar_rows=sorted(ar_rows,key=lambda x:x['age'],reverse=True), ap_rows=sorted(ap_rows,key=lambda x:x['age'],reverse=True), today=today,
        ar_totals=[sum(r['buckets'][i] for r in ar_rows) for i in range(5)], ap_totals=[sum(r['buckets'][i] for r in ap_rows) for i in range(5)])

@app.route('/reports/aging/export.xlsx')
@login_required
def export_aging_report():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    sid=store_id(); today=date.today(); wb=Workbook();
    for title, records, is_ar in [('Aged Receivables', Sale.query.filter_by(supermarket_id=sid).all(), True), ('Aged Payables', Purchase.query.filter_by(supermarket_id=sid).all(), False)]:
        ws=wb.active if title=='Aged Receivables' else wb.create_sheet()
        ws.title=title
        ws.append(['Invoice / Reference','Invoice Date','Customer / Supplier','Total','Paid','Returns','Outstanding','Age (days)','Current','1-30 days','31-60 days','61-90 days','Over 90 days'])
        for obj in records:
            if is_ar:
                returned=db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(sale_id=obj.id,supermarket_id=sid).scalar() or 0; due=max(0,obj.total-(obj.paid or 0)-returned); dt=obj.sale_date.date(); ref=obj.invoice_no; party=obj.customer_name
            else:
                returned=db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(purchase_id=obj.id,supermarket_id=sid).scalar() or 0; due=max(0,obj.total-(obj.paid or 0)-returned); dt=obj.purchase_date; ref=f'PUR-{obj.id}'; party=obj.supplier or 'Supplier not specified'
            if due <= .009: continue
            age=(today-dt).days; b=[due if age<=0 else 0,due if 1<=age<=30 else 0,due if 31<=age<=60 else 0,due if 61<=age<=90 else 0,due if age>90 else 0]
            ws.append([ref,dt.isoformat(),party,obj.total,obj.paid or 0,returned,round(due,2),age,*b])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='203B53')
        for col in ws.columns: ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),30)
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f'supermart_aging_{today.isoformat()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/settings/lists', methods=['GET','POST'])
@login_required
def settings_lists():
    sid=store_id()
    if request.method == 'POST':
        kind=request.form.get('kind',''); name=request.form.get('name','').strip()
        if not name:
            flash('Enter a name.', 'error')
        elif kind == 'category':
            if ExpenseCategory.query.filter_by(supermarket_id=sid,name=name).first(): flash('That expense category already exists.', 'error')
            else: db.session.add(ExpenseCategory(supermarket_id=sid,name=name)); db.session.commit(); flash('Expense category added.', 'success')
        elif kind == 'payment':
            if PaymentMethod.query.filter_by(supermarket_id=sid,name=name).first(): flash('That payment method already exists.', 'error')
            else: db.session.add(PaymentMethod(supermarket_id=sid,name=name)); db.session.commit(); flash('Payment method added.', 'success')
        else: flash('Invalid list type.', 'error')
        return redirect(url_for('settings_lists'))
    return render_template('settings_lists.html', categories=ExpenseCategory.query.filter_by(supermarket_id=sid).order_by(ExpenseCategory.name).all(), methods=PaymentMethod.query.filter_by(supermarket_id=sid).order_by(PaymentMethod.name).all())

@app.route('/settings/lists/<kind>/<int:item_id>/toggle', methods=['POST'])
@login_required
def toggle_setting_list(kind,item_id):
    sid=store_id()
    model=ExpenseCategory if kind=='category' else PaymentMethod if kind=='payment' else None
    if model is None: return ('Not found',404)
    item=model.query.filter_by(id=item_id,supermarket_id=sid).first_or_404(); item.active=not item.active; db.session.commit()
    flash('Setting updated.', 'success'); return redirect(url_for('settings_lists'))

@app.route('/reports/suppliers')
@login_required
def supplier_report():
    sid=store_id(); start,end=date_range()
    purchases=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all()
    grouped={}
    for purchase in purchases:
        name=(purchase.supplier or 'Unspecified supplier').strip()
        row=grouped.setdefault(name, {'name':name,'purchases':0.0,'returns':0.0,'paid':0.0,'due':0.0,'count':0})
        returned=db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(supermarket_id=sid,purchase_id=purchase.id).scalar() or 0
        row['purchases'] += purchase.total; row['returns'] += returned; row['paid'] += purchase.paid or 0
        row['due'] += max(0,purchase.total-returned-(purchase.paid or 0)); row['count'] += 1
    rows=sorted(grouped.values(),key=lambda x:x['name'].lower())
    totals={key:sum(row[key] for row in rows) for key in ('purchases','returns','paid','due','count')}
    return render_template('supplier_report.html', rows=rows,totals=totals,start=start.isoformat(),end=end.isoformat(),title='Supplier / Vendor-wise Total Report')

@app.route('/reports/customers')
@login_required
def customer_report():
    sid=store_id(); start,end=date_range()
    sales=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start,func.date(Sale.sale_date)<=end).all()
    grouped={}
    for sale in sales:
        name=(sale.customer_name or 'Walk-in customer').strip() or 'Walk-in customer'
        row=grouped.setdefault(name, {'name':name,'sales':0.0,'returns':0.0,'received':0.0,'due':0.0,'count':0})
        returned=db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(supermarket_id=sid,sale_id=sale.id).scalar() or 0
        row['sales'] += sale.total or 0
        row['returns'] += returned
        row['received'] += sale.paid or 0
        row['due'] += max(0,(sale.total or 0)-returned-(sale.paid or 0))
        row['count'] += 1
    rows=sorted(grouped.values(),key=lambda x:x['name'].lower())
    totals={key:sum(row[key] for row in rows) for key in ('sales','returns','received','due','count')}
    return render_template('customer_report.html',rows=rows,totals=totals,start=start.isoformat(),end=end.isoformat(),title='Customer-wise Total Report')

@app.route('/reports/customers/export.xlsx')
@login_required
def customer_report_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    sid=store_id(); start,end=date_range()
    sales=Sale.query.filter(Sale.supermarket_id==sid,func.date(Sale.sale_date)>=start,func.date(Sale.sale_date)<=end).all()
    grouped={}
    for sale in sales:
        name=(sale.customer_name or 'Walk-in customer').strip() or 'Walk-in customer'
        row=grouped.setdefault(name,[name,0.0,0.0,0.0,0.0,0])
        returned=db.session.query(func.coalesce(func.sum(SalesReturn.amount),0)).filter_by(supermarket_id=sid,sale_id=sale.id).scalar() or 0
        row[1]+=sale.total or 0; row[2]+=returned; row[3]+=sale.paid or 0; row[4]+=max(0,(sale.total or 0)-returned-(sale.paid or 0)); row[5]+=1
    wb=Workbook(); ws=wb.active; ws.title='Customer Report'
    ws.append(['Customer','Sales Total','Sales Returns','Received','Outstanding','Invoices'])
    ordered=sorted(grouped.values(),key=lambda x:x[0].lower())
    for row in ordered: ws.append(row)
    ws.append(['GRAND TOTAL']+[sum(row[i] for row in ordered) for i in range(1,6)])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='203B53')
    for c in ws[ws.max_row]: c.font=Font(bold=True)
    for col in ws.columns: ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f'customer_report_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/reports/suppliers/export.xlsx')
@login_required
def supplier_report_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    sid=store_id(); start,end=date_range(); purchases=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.purchase_date>=start,Purchase.purchase_date<=end).all(); grouped={}
    for p in purchases:
        name=(p.supplier or 'Unspecified supplier').strip(); row=grouped.setdefault(name,[name,0.0,0.0,0.0,0.0,0]); ret=db.session.query(func.coalesce(func.sum(PurchaseReturn.amount),0)).filter_by(supermarket_id=sid,purchase_id=p.id).scalar() or 0
        row[1]+=p.total; row[2]+=ret; row[3]+=p.paid or 0; row[4]+=max(0,p.total-ret-(p.paid or 0)); row[5]+=1
    wb=Workbook(); ws=wb.active; ws.title='Supplier Vendor Report'; ws.append(['Supplier / Vendor','Purchase Total','Purchase Returns','Paid','Outstanding','Invoices'])
    ordered=sorted(grouped.values(),key=lambda x:x[0].lower())
    for row in ordered: ws.append(row)
    ws.append(['GRAND TOTAL']+[sum(row[i] for row in ordered) for i in range(1,6)])
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='203B53')
    for c in ws[ws.max_row]: c.font=Font(bold=True)
    for col in ws.columns: ws.column_dimensions[get_column_letter(col[0].column)].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
    out=BytesIO(); wb.save(out); out.seek(0); return send_file(out,as_attachment=True,download_name=f'supplier_vendor_report_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/reports/expiry')
@login_required
def expiry_report():
    sid=store_id(); today=date.today(); days=max(1,min(180,int(request.args.get('days',30) or 30))); cutoff=today+timedelta(days=days)
    products=Product.query.filter(Product.supermarket_id==sid,Product.active.is_(True),Product.expiry_date.isnot(None),Product.expiry_date<=cutoff).order_by(Product.expiry_date.asc()).all()
    batches=Purchase.query.filter(Purchase.supermarket_id==sid,Purchase.expiry_date.isnot(None),Purchase.expiry_date<=cutoff).order_by(Purchase.expiry_date.asc()).all()
    return render_template('expiry_report.html',products=products,batches=batches,today=today,cutoff=cutoff,days=days)

@app.route('/reports/expiry/export.xlsx')
@login_required
def expiry_report_export():
    from openpyxl import Workbook
    sid=store_id(); today=date.today(); days=max(1,min(180,int(request.args.get('days',30) or 30))); cutoff=today+timedelta(days=days)
    products=Product.query.filter(Product.supermarket_id==sid,Product.active.is_(True),Product.expiry_date.isnot(None),Product.expiry_date<=cutoff).order_by(Product.expiry_date.asc()).all()
    wb=Workbook(); ws=wb.active; ws.title='Expiry Alert'; ws.append(['Product','SKU','Stock','Expiry date','Days remaining','Status'])
    for p in products:
        remaining=(p.expiry_date-today).days; ws.append([p.name,p.sku or '',p.stock_qty,p.expiry_date.isoformat(),remaining,'Expired' if remaining<0 else 'Near expiry'])
    out=BytesIO(); wb.save(out); out.seek(0); return send_file(out,as_attachment=True,download_name=f'expiry_alert_{today}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/health')
def health(): return {'status':'ok'}

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=os.environ.get('FLASK_DEBUG')=='1')
